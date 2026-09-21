"""Đọc giáo trình, slide, đề thi thành chunk.

Quy ước thư mục: ``data/raw/<MÃ MÔN>/<file>`` — môn học suy ra từ tên thư mục,
nên chỉ cần thả file đúng chỗ là xong, không phải khai báo gì thêm.

Cảnh báo thật về chất lượng đầu vào, cần nêu trong phần Limitations của báo cáo:
PDF sinh từ LaTeX thường cho ra công thức bị vỡ khi trích text (``∫`` thành
khoảng trắng, chỉ số trên/dưới mất vị trí). Phần VĂN XUÔI vẫn dùng tốt cho
truy xuất, nhưng đừng trông đợi trích được công thức nguyên vẹn. Với các công
thức quan trọng, gõ tay vào file ``pitfalls`` sẽ cho chất lượng cao hơn nhiều
so với mọi thủ thuật trích tự động.
"""

from __future__ import annotations

import re
from pathlib import Path

from src.rag.chunking import Chunk, ChunkType, assign_ids, chunk_text

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".pptx", ".txt", ".md"}

#: Ký tự rác hay gặp khi trích text từ PDF sinh bởi LaTeX.
_CLEANUP = (
    (re.compile(r"[ \t]{2,}"), " "),
    (re.compile(r"\n{3,}"), "\n\n"),
    (re.compile(r"(?<=\w)-\n(?=\w)"), ""),   # nối từ bị gạch nối xuống dòng
    (re.compile(r"[ﬀ-ﬆ]"), ""),     # ligature ff/fi/fl trích ra bị lỗi
)


def clean_text(text: str) -> str:
    for pattern, replacement in _CLEANUP:
        text = pattern.sub(replacement, text)
    return text.strip()


# --------------------------------------------------------------------------
# Đọc từng định dạng
# --------------------------------------------------------------------------

def read_pdf(path: Path) -> list[tuple[int, str]]:
    """Trả về danh sách ``(số trang, nội dung)``.

    Giữ số trang để trích dẫn nguồn được chính xác — sinh viên cần tra lại
    được trong sách, nếu không thì phần trích dẫn chỉ là trang trí.
    """
    import fitz

    pages: list[tuple[int, str]] = []
    with fitz.open(path) as document:
        for number, page in enumerate(document, start=1):
            content = clean_text(page.get_text())
            if content:
                pages.append((number, content))
    return pages


def read_docx(path: Path) -> list[tuple[int, str]]:
    import docx

    document = docx.Document(str(path))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))

    return [(1, clean_text("\n\n".join(paragraphs)))] if paragraphs else []


def read_pptx(path: Path) -> list[tuple[int, str]]:
    try:
        from pptx import Presentation
    except ImportError:
        # python-pptx là tuỳ chọn. Thiếu nó thì bỏ qua slide chứ không làm
        # hỏng cả lần nạp — các định dạng khác vẫn phải vào được.
        return []

    slides: list[tuple[int, str]] = []
    presentation = Presentation(str(path))
    for number, slide in enumerate(presentation.slides, start=1):
        texts = [
            shape.text for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text.strip()
        ]
        if texts:
            slides.append((number, clean_text("\n".join(texts))))
    return slides


def read_plain(path: Path) -> list[tuple[int, str]]:
    return [(1, clean_text(path.read_text(encoding="utf-8", errors="replace")))]


_READERS = {
    ".pdf": read_pdf,
    ".docx": read_docx,
    ".pptx": read_pptx,
    ".txt": read_plain,
    ".md": read_plain,
}


def read_document(path: Path) -> list[tuple[int, str]]:
    reader = _READERS.get(path.suffix.lower())
    if reader is None:
        return []
    try:
        return reader(path)
    except Exception as exc:
        # Một file hỏng không được làm chết cả lần nạp: giáo trình thật hay có
        # file bị khoá, PDF lỗi, hoặc định dạng lạ.
        print(f"  [bỏ qua] {path.name}: {type(exc).__name__}: {exc}")
        return []


# --------------------------------------------------------------------------
# Nạp cả thư mục
# --------------------------------------------------------------------------

def infer_chapter(path: Path, subject_root: Path) -> str:
    """Suy ra tên chương từ đường dẫn tương đối."""
    relative = path.relative_to(subject_root)
    if len(relative.parts) > 1:
        return relative.parts[0]
    return path.stem


def ingest_file(path: Path, subject: str, subject_root: Path) -> list[Chunk]:
    chapter = infer_chapter(path, subject_root)
    chunks: list[Chunk] = []
    for page, content in read_document(path):
        chunks.extend(chunk_text(
            content, subject=subject, chapter=chapter, source=path.name, page=page
        ))
    return chunks


def ingest_subject(subject_root: Path, subject: str | None = None) -> list[Chunk]:
    """Nạp toàn bộ tài liệu của một môn."""
    subject = subject or subject_root.name
    chunks: list[Chunk] = []

    for path in sorted(subject_root.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES:
            chunks.extend(ingest_file(path, subject, subject_root))

    return chunks


def ingest_all(raw_root: Path, subjects: list[str] | None = None) -> list[Chunk]:
    """Nạp mọi môn trong ``data/raw/``, mỗi thư mục con là một môn."""
    chunks: list[Chunk] = []
    for subject_root in sorted(raw_root.iterdir()):
        if not subject_root.is_dir():
            continue
        if subjects and subject_root.name not in subjects:
            continue
        found = ingest_subject(subject_root)
        print(f"  {subject_root.name}: {len(found)} chunk")
        chunks.extend(found)
    return assign_ids(chunks)


def load_pitfalls(path: Path, subject: str) -> list[Chunk]:
    """Nạp bảng 'bẫy thường gặp' viết tay từ taxonomy lỗi.

    Đây là nguồn chất lượng cao nhất của cả kho: viết tay, ngắn gọn, gắn đúng
    mã lỗi. Hint mức 2 lấy trực tiếp từ đây, nên đừng trộn lẫn với văn xuôi
    trích tự động từ PDF.
    """
    import json

    entries = json.loads(path.read_text(encoding="utf-8"))
    return [
        Chunk(
            text=f"{entry['name_vi']}. {entry.get('description_vi', '')} {entry.get('hint_vi', '')}".strip(),
            chunk_type=ChunkType.PITFALL,
            subject=subject,
            chapter="Bẫy thường gặp",
            skill_ids=[entry["skill_id"]] if entry.get("skill_id") else [],
            source=f"taxonomy:{entry['code']}",
        )
        for entry in entries
    ]
